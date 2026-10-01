"""Turn configured sources into chunks: {id, text, title, source, section, page, meta}.

Source types (config.yaml -> collections.<name>.sources):
  wiki_dir     Markdown wiki export: <area>/<slug>.md; an optional line "Page: `/path`" (or "Page in <App>: `/path`") is kept as the page link
  markdown_dir any folder of .md/.txt files
  qa_jsonl     question/answer library, one JSON object per line: {requirement|question, answer, compliance?, notes?, module?, source?, year?, section?, id?}
  files_dir    drop folder: .md .txt .pdf .docx .xlsx (and nested folders)
"""
import glob
import hashlib
import json
import logging
import os
import re

log = logging.getLogger("rag.ingest")
FILE_TYPES = (".md", ".txt", ".pdf", ".docx", ".xlsx")


# ----------------------------------------------------------------- source fingerprint

def source_files(src):
    p = src["path"]
    if src["type"] in ("qa_jsonl", "rfp_jsonl"):
        return [p] if os.path.exists(p) else []
    exts = (".md", ".txt") if src["type"] in ("wiki_dir", "markdown_dir") else FILE_TYPES
    return sorted(f for f in glob.glob(os.path.join(p, "**", "*"), recursive=True)
                  if os.path.isfile(f) and f.lower().endswith(exts) and not os.path.basename(f).startswith(("._", "~$")))


def fingerprint(sources, extra=""):
    h = hashlib.sha256(extra.encode())
    for src in sources:
        for f in source_files(src):
            st = os.stat(f)
            h.update(f"{f}|{st.st_size}|{int(st.st_mtime)}".encode())
    return h.hexdigest()


# ----------------------------------------------------------------- chunking

def split_markdown(text, max_chars, overlap):
    """Split on headings, then pack paragraphs up to max_chars. Returns [(section, text)]."""
    sections, cur_head, buf = [], "", []
    for line in text.splitlines():
        if re.match(r"^#{1,4}\s", line):
            if buf:
                sections.append((cur_head, "\n".join(buf).strip()))
            cur_head, buf = line.lstrip("#").strip(), []
        else:
            buf.append(line)
    if buf:
        sections.append((cur_head, "\n".join(buf).strip()))
    out = []
    for head, body in sections:
        if not body:
            continue
        paras = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
        chunk = ""
        for p in paras:
            while len(p) > max_chars:  # very long paragraph: hard split
                if chunk:
                    out.append((head, chunk)); chunk = ""
                out.append((head, p[:max_chars])); p = p[max_chars - overlap:]
            if chunk and len(chunk) + len(p) + 2 > max_chars:
                out.append((head, chunk))
                chunk = chunk[-overlap:] + "\n\n" + p if overlap else p
            else:
                chunk = f"{chunk}\n\n{p}" if chunk else p
        if chunk:
            out.append((head, chunk))
    return out


def _cid(*parts):
    return hashlib.sha1("|".join(map(str, parts)).encode()).hexdigest()[:16]


def _mk(title, section, text, source, page="", meta=None):
    header = title + (f" — {section}" if section and section != title else "")
    return {"id": _cid(source, section, text[:200]), "title": title, "section": section or "", "page": page or "",
            "source": source, "text": f"{header}\n{text}", "meta": json.dumps(meta or {}, ensure_ascii=False)}


# ----------------------------------------------------------------- loaders

def load_wiki(src, s):
    for f in source_files(src):
        raw = open(f, encoding="utf-8", errors="replace").read()
        title = (re.search(r"^#\s+(.+)$", raw, re.M) or [None, os.path.basename(f)[:-3]])[1].strip()
        page = (re.search(r"Page(?: in [^:\n]+)?:\s*`([^`]+)`", raw) or [None, ""])[1]
        area = os.path.basename(os.path.dirname(f))
        body = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", raw)          # drop image refs
        body = re.sub(r"^#\s+.+$", "", body, count=1, flags=re.M)
        for sec, txt in split_markdown(body, s.chunk_chars, s.chunk_overlap):
            yield _mk(title, sec, txt, f"wiki/{area}/{os.path.basename(f)}", page, {"area": area})


def load_markdown_dir(src, s):
    base = src["path"]
    for f in source_files(src):
        raw = open(f, encoding="utf-8", errors="replace").read()
        rel = os.path.relpath(f, base)
        title = (re.search(r"^#\s+(.+)$", raw, re.M) or [None, os.path.splitext(os.path.basename(f))[0]])[1].strip()
        if os.path.basename(f) in ("INDEX.md",):
            continue
        for sec, txt in split_markdown(raw, s.chunk_chars, s.chunk_overlap):
            yield _mk(title, sec, txt, f"{src.get('label', 'docs')}/{rel}", meta={"folder": os.path.dirname(rel)})


def load_qa_jsonl(src, s):
    for n, line in enumerate(open(src["path"], encoding="utf-8"), 1):
        if not line.strip():
            continue
        e = json.loads(line)
        e.setdefault("requirement", e.get("question", ""))
        e.setdefault("id", f"QA-{n:04d}")
        e.setdefault("source", src.get("label", "Q&A library"))
        e.setdefault("year", "")
        parts = [f"Requirement: {e['requirement']}"]
        if e.get("compliance"):
            parts.append(f"Compliance declared: {e['compliance']}")
        if e.get("internal_verdict"):
            parts.append(f"Internal verdict: {e['internal_verdict']}")
        for k, lab in (("module", "Product module"), ("answer", "Answer given"), ("notes", "Internal notes")):
            if e.get(k):
                parts.append(f"{lab}: {e[k]}")
        title = f"Past answer ({e['source']}{', ' + str(e['year']) if e.get('year') else ''})"
        meta = {k: e.get(k) for k in ("id", "source", "year", "lang", "compliance", "source_ref")}
        yield _mk(title, e.get("section", ""), "\n".join(parts), f"{src.get('label', 'qa')}/{e['id']}", meta=meta)


def _pdf_text(f):
    import fitz  # pymupdf
    with fitz.open(f) as d:
        return "\n\n".join(p.get_text() for p in d)


def _docx_text(f):
    import docx
    d = docx.Document(f)
    out = []
    for block in d.element.body.iterchildren():
        tag = block.tag.split("}")[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            p = Paragraph(block, d)
            if p.text.strip():
                style = (p.style.name if p.style is not None else "") or ""
                out.append(("## " if style.lower().startswith(("heading", "title")) else "") + p.text.strip())
        elif tag == "tbl":
            from docx.table import Table
            for row in Table(block, d).rows:
                cells, seen = [], set()
                for c in row.cells:
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc)); cells.append(c.text.strip())
                if any(cells):
                    out.append(" | ".join(cells))
    return "\n\n".join(out)


def _xlsx_chunks(f, s):
    import openpyxl
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [[("" if c is None else str(c).strip()) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        hdr_i = next((i for i, r in enumerate(rows[:20]) if sum(1 for c in r if c and len(c) < 60) >= 2), 0)
        hdr = rows[hdr_i]
        buf = ""
        for r in rows[hdr_i + 1:]:
            line = "; ".join(f"{hdr[i] if i < len(hdr) and hdr[i] else f'col{i + 1}'}: {v}" for i, v in enumerate(r) if v)
            if buf and len(buf) + len(line) > s.chunk_chars:
                yield ws.title, buf; buf = ""
            buf = f"{buf}\n{line}" if buf else line
        if buf:
            yield ws.title, buf


def load_files_dir(src, s):
    base = src["path"]
    for f in source_files(src):
        rel, title = os.path.relpath(f, base), os.path.splitext(os.path.basename(f))[0]
        label = f"{src.get('label', 'files')}/{rel}"
        try:
            ext = os.path.splitext(f)[1].lower()
            if ext == ".xlsx":
                for sheet, txt in _xlsx_chunks(f, s):
                    yield _mk(title, sheet, txt, label)
                continue
            text = (_pdf_text(f) if ext == ".pdf" else _docx_text(f) if ext == ".docx"
                    else open(f, encoding="utf-8", errors="replace").read())
            for sec, txt in split_markdown(text, s.chunk_chars, s.chunk_overlap):
                yield _mk(title, sec, txt, label)
        except Exception as e:  # one bad file must not stop the collection
            log.warning("skip %s: %s", f, e)


LOADERS = {"wiki_dir": load_wiki, "markdown_dir": load_markdown_dir, "qa_jsonl": load_qa_jsonl, "rfp_jsonl": load_qa_jsonl,
           "files_dir": load_files_dir}


def anonymiser(names):
    if not names:
        return lambda t: t
    rx = re.compile("|".join(names), re.I)
    return lambda t: rx.sub("the customer", t)


def build_chunks(col, s):
    scrub = anonymiser(s.customer_names) if col.anonymise else (lambda t: t)
    seen, out = set(), []
    for src in col.sources:
        loader = LOADERS.get(src["type"])
        if not loader:
            log.warning("unknown source type %s", src["type"]); continue
        for ch in loader(src, s):
            if len(ch["text"]) < 40 or ch["id"] in seen:
                continue
            seen.add(ch["id"])
            ch["text"] = scrub(ch["text"]); ch["title"] = scrub(ch["title"])
            out.append(ch)
    return out
